import os
import re
import json
import datetime as _dt
from decimal import Decimal

from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse
from django.db.models import Q
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.utils import timezone

from ..models import (
    TPAScheme, TPAPatient, TPADocument,
    TPAImportBatch, TPAImportRow, TPAImportMapping,
    TPA_STATUS_CHOICES, TPA_SCHEME_TYPE_CHOICES, TPA_DOCUMENT_TYPE_CHOICES,
)


TPA_IMPORT_TARGET_FIELDS = [
    ('tid', 'TID / Pre-auth Number'),
    ('policy_no', 'Policy No. / Beneficiary ID'),
    ('patient_name', 'Patient Name'),
    ('package_code', 'Diagnosis / Package Code'),
    ('status', 'Case Status'),
    ('approved_amount', 'Approved Amount'),
    ('admission_date', 'Admission Date'),
    ('scheme_name', 'TPA / Scheme Name'),
]


TPA_IMPORT_GUESS_KEYWORDS = {
    'tid': ['tid'],
    'policy_no': ['policy', 'beneficiary'],
    'patient_name': ['patient name', 'patient_name', 'name'],
    'package_code': ['pkg code', 'package code', 'pkg_code', 'diagnos'],
    'status': ['case status', 'status'],
    'approved_amount': ['amount'],
    'admission_date': ['admission', 'admit date'],
    'scheme_name': ['scheme', 'tpa name', 'insurance'],
}


@login_required
def tpa_patient_list(request):
    scheme_type = request.GET.get('scheme_type', '')
    status = request.GET.get('status', '')
    q = request.GET.get('q', '').strip()

    cases = TPAPatient.objects.select_related('patient', 'scheme').all()
    if scheme_type:
        cases = cases.filter(scheme_type=scheme_type)
    if status:
        cases = cases.filter(status=status)
    if q:
        cases = cases.filter(
            Q(patient__full_name__icontains=q) | Q(patient__uhid__icontains=q) |
            Q(patient_name_raw__icontains=q) | Q(tid_number__icontains=q) |
            Q(policy_no__icontains=q)
        )

    return render(request, 'tpa/list.html', {
        'cases': cases,
        'scheme_type': scheme_type,
        'status': status,
        'q': q,
        'status_choices': TPA_STATUS_CHOICES,
        'draft_count': TPAPatient.objects.filter(is_draft=True).count(),
    })


@login_required
def tpa_patient_create(request):
    schemes = TPAScheme.objects.filter(is_active=True).order_by('name')

    if request.method == 'POST':
        scheme_type = request.POST.get('scheme_type', '')
        if not scheme_type:
            messages.error(request, 'Scheme type is required.')
        else:
            tpa = TPAPatient.objects.create(
                patient_id=request.POST.get('patient_id') or None,
                scheme_type=scheme_type,
                scheme_id=request.POST.get('scheme') or None,
                tid_number=request.POST.get('tid_number', '').strip(),
                policy_no=request.POST.get('policy_no', '').strip(),
                admission_date=request.POST.get('admission_date') or None,
                diagnosis=request.POST.get('diagnosis', '').strip(),
                package_code=request.POST.get('package_code', '').strip(),
                approved_amount=request.POST.get('approved_amount') or None,
                status=request.POST.get('status', 'pending'),
                remarks=request.POST.get('remarks', '').strip(),
                created_by=request.user,
            )
            for f in request.FILES.getlist('preauth_letter'):
                TPADocument.objects.create(tpa_patient=tpa, doc_type='preauth_letter', file=f, uploaded_by=request.user)
            for f in request.FILES.getlist('discharge_summary'):
                TPADocument.objects.create(tpa_patient=tpa, doc_type='discharge_summary', file=f, uploaded_by=request.user)
            messages.success(request, 'TPA case created.')
            return redirect('hms:tpa_patient_detail', pk=tpa.pk)

    return render(request, 'tpa/form.html', {
        'action': 'Add',
        'schemes': schemes,
        'scheme_type_pre': request.GET.get('scheme_type', 'private'),
        'status_choices': TPA_STATUS_CHOICES,
        'scheme_type_choices': TPA_SCHEME_TYPE_CHOICES,
    })


@login_required
def tpa_patient_edit(request, pk):
    tpa = get_object_or_404(TPAPatient, pk=pk)
    schemes = TPAScheme.objects.filter(is_active=True).order_by('name')

    if request.method == 'POST':
        tpa.patient_id = request.POST.get('patient_id') or None
        tpa.scheme_type = request.POST.get('scheme_type', tpa.scheme_type)
        tpa.scheme_id = request.POST.get('scheme') or None
        tpa.tid_number = request.POST.get('tid_number', '').strip()
        tpa.policy_no = request.POST.get('policy_no', '').strip()
        tpa.admission_date = request.POST.get('admission_date') or None
        tpa.diagnosis = request.POST.get('diagnosis', '').strip()
        tpa.package_code = request.POST.get('package_code', '').strip()
        tpa.approved_amount = request.POST.get('approved_amount') or None
        tpa.status = request.POST.get('status', 'pending')
        tpa.remarks = request.POST.get('remarks', '').strip()
        if tpa.patient_id:
            tpa.is_draft = False
        tpa.save()

        for f in request.FILES.getlist('preauth_letter'):
            TPADocument.objects.create(tpa_patient=tpa, doc_type='preauth_letter', file=f, uploaded_by=request.user)
        for f in request.FILES.getlist('discharge_summary'):
            TPADocument.objects.create(tpa_patient=tpa, doc_type='discharge_summary', file=f, uploaded_by=request.user)

        messages.success(request, 'TPA case updated.')
        return redirect('hms:tpa_patient_detail', pk=tpa.pk)

    return render(request, 'tpa/form.html', {
        'action': 'Edit',
        'tpa': tpa,
        'schemes': schemes,
        'status_choices': TPA_STATUS_CHOICES,
        'scheme_type_choices': TPA_SCHEME_TYPE_CHOICES,
    })


@login_required
def tpa_patient_detail(request, pk):
    tpa = get_object_or_404(TPAPatient.objects.select_related('patient', 'scheme'), pk=pk)
    documents = tpa.documents.all()
    import_rows = tpa.import_rows.select_related('batch').order_by('-id')
    return render(request, 'tpa/detail.html', {
        'tpa': tpa,
        'documents': documents,
        'import_rows': import_rows,
        'doc_type_choices': TPA_DOCUMENT_TYPE_CHOICES,
    })


@login_required
def tpa_patient_delete(request, pk):
    tpa = get_object_or_404(TPAPatient, pk=pk)
    if request.method == 'POST':
        tpa.delete()
        messages.success(request, 'TPA case deleted.')
        return redirect('hms:tpa_patient_list')
    return redirect('hms:tpa_patient_detail', pk=pk)


@login_required
def tpa_document_upload(request, pk):
    tpa = get_object_or_404(TPAPatient, pk=pk)
    if request.method == 'POST':
        f = request.FILES.get('file')
        if f:
            TPADocument.objects.create(
                tpa_patient=tpa,
                doc_type=request.POST.get('doc_type', 'other'),
                file=f,
                uploaded_by=request.user,
            )
            messages.success(request, 'Document uploaded.')
        else:
            messages.error(request, 'Choose a file to upload.')
    return redirect('hms:tpa_patient_detail', pk=pk)


@login_required
def tpa_document_delete(request, doc_id):
    doc = get_object_or_404(TPADocument, id=doc_id)
    tpa_id = doc.tpa_patient_id
    if request.method == 'POST':
        if doc.file:
            try:
                if os.path.isfile(doc.file.path):
                    os.remove(doc.file.path)
            except Exception:
                pass
        doc.delete()
        messages.success(request, 'Document removed.')
    return redirect('hms:tpa_patient_detail', pk=tpa_id)


@login_required
def tpa_scheme_add(request):
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)
    data = json.loads(request.body)
    name = data.get('name', '').strip()
    scheme_type = data.get('scheme_type', 'private')
    if not name:
        return JsonResponse({'error': 'Name required'}, status=400)
    # scheme_type is part of the lookup (not just a default) so the same name
    # under a different category creates its own row instead of silently
    # reusing — and locking into — whichever category first created it.
    scheme, created = TPAScheme.objects.get_or_create(name=name, scheme_type=scheme_type)
    return JsonResponse({'id': scheme.id, 'name': scheme.name, 'scheme_type': scheme.scheme_type, 'created': created})


@login_required
def tpa_scheme_search(request):
    q = request.GET.get('q', '').strip()
    scheme_type = request.GET.get('scheme_type', '')
    schemes = TPAScheme.objects.filter(is_active=True)
    if scheme_type:
        schemes = schemes.filter(scheme_type=scheme_type)
    if q:
        schemes = schemes.filter(name__icontains=q)
    data = list(schemes.order_by('name').values('id', 'name', 'scheme_type')[:10])
    return JsonResponse({'schemes': data})


def _tpa_read_headers(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    row = next(ws.iter_rows(values_only=True))
    wb.close()
    return [str(c).strip() if c not in (None, '') else f'Column {i + 1}' for i, c in enumerate(row)]


def _tpa_guess_mapping(headers):
    guessed = {}
    for target, keywords in TPA_IMPORT_GUESS_KEYWORDS.items():
        for h in headers:
            hl = h.lower()
            if any(kw in hl for kw in keywords):
                guessed[target] = h
                break
    return guessed


def _tpa_json_safe(v):
    if v is None:
        return None
    if isinstance(v, (_dt.date, _dt.datetime)):
        return v.isoformat()
    if isinstance(v, Decimal):
        return str(v)
    return v


def _tpa_parse_amount(raw):
    if raw is None or raw == '':
        return None
    if isinstance(raw, (int, float, Decimal)):
        return Decimal(str(raw))
    s = re.sub(r'[^0-9.\-]', '', str(raw))
    if not s:
        return None
    try:
        return Decimal(s)
    except Exception:
        return None


def _tpa_parse_date(raw):
    if raw is None or raw == '':
        return None
    if isinstance(raw, _dt.datetime):
        return raw.date()
    if isinstance(raw, _dt.date):
        return raw
    s = str(raw).strip()
    for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d.%m.%Y'):
        try:
            return _dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    # Government-scheme portal exports (e.g. MAA Yojana / Ayushman Bharat
    # "Generic Search Report") write dates as "02,August   , 2026" — comma
    # separated with irregular internal whitespace, sometimes with a
    # trailing " 12:00 AM". Strip that off and re-join on single spaces.
    s_no_time = re.sub(r'\s+\d{1,2}:\d{2}\s*[AP]M\s*$', '', s, flags=re.IGNORECASE)
    parts = [p.strip() for p in s_no_time.split(',') if p.strip()]
    if len(parts) == 3:
        joined = ' '.join(parts)
        for fmt in ('%d %B %Y', '%d %b %Y'):
            try:
                return _dt.datetime.strptime(joined, fmt).date()
            except ValueError:
                continue
    return None


def _tpa_normalize_status(raw):
    if not raw:
        return None
    s = str(raw).strip().lower()
    for kw, val in [('approv', 'approved'), ('settl', 'settled'), ('reject', 'rejected'),
                     ('query', 'query_raised'), ('paid', 'settled'), ('pend', 'pending')]:
        if kw in s:
            return val
    return None


def _run_tpa_import(batch):
    import openpyxl
    wb = openpyxl.load_workbook(batch.file.path, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        batch.status = 'failed'
        batch.error_message = 'The uploaded file has no rows.'
        batch.save(update_fields=['status', 'error_message'])
        return

    headers = [str(c).strip() if c not in (None, '') else f'Column {i + 1}' for i, c in enumerate(header_row)]
    mapping = batch.mapping or {}

    def cell(row_dict, target):
        h = mapping.get(target)
        if not h:
            return None
        return row_dict.get(h)

    # First pass: some source reports (e.g. government-scheme "Generic Search
    # Report" exports) list one row per claim LINE ITEM (main procedure,
    # implant/consumable add-ons, ...) all sharing the same TID rather than
    # one row per claim. Pre-sum the approved amount across every row that
    # shares a TID so the total isn't silently overwritten down to just the
    # last line's amount.
    parsed_rows = []
    row_num = 1
    for raw_row in rows_iter:
        row_num += 1
        if raw_row is None or all(c is None or str(c).strip() == '' for c in raw_row):
            continue
        row_dict = dict(zip(headers, raw_row))
        tid_val = str(cell(row_dict, 'tid') or '').strip()
        parsed_rows.append({
            'row_num': row_num,
            'row_dict': row_dict,
            'raw_data': {k: _tpa_json_safe(v) for k, v in row_dict.items()},
            'tid_val': tid_val,
            'policy_val': str(cell(row_dict, 'policy_no') or '').strip(),
            'name_val': str(cell(row_dict, 'patient_name') or '').strip(),
        })

    tid_totals = {}
    for pr in parsed_rows:
        if pr['tid_val']:
            amt = _tpa_parse_amount(cell(pr['row_dict'], 'approved_amount')) or Decimal('0')
            tid_totals[pr['tid_val']] = tid_totals.get(pr['tid_val'], Decimal('0')) + amt

    updated = created = merged = skipped = 0
    tid_seen = {}  # tid_val -> the TPAPatient object handling that TID's total, once processed
    for pr in parsed_rows:
        row_num = pr['row_num']
        row_dict = pr['row_dict']
        raw_data = pr['raw_data']
        tid_val = pr['tid_val']
        policy_val = pr['policy_val']
        name_val = pr['name_val']

        if not tid_val and not policy_val and not name_val:
            TPAImportRow.objects.create(
                batch=batch, row_number=row_num, raw_data=raw_data,
                match_type='none', action='skipped',
                error_message='No TID, Policy No. or Patient Name in this row.',
            )
            skipped += 1
            continue

        if tid_val and tid_val in tid_seen:
            # A later line item for a TID already processed above — its
            # amount is already folded into that row's total, so this row
            # just records the audit trail, no further DB change.
            TPAImportRow.objects.create(
                batch=batch, row_number=row_num, raw_data=raw_data,
                match_type='tid', tpa_patient=tid_seen[tid_val], action='merged',
                error_message=(
                    f"Line item amount merged into TID {tid_val}'s total "
                    f"(₹{tid_totals[tid_val]})."
                ),
            )
            merged += 1
            continue

        qs = TPAPatient.objects.filter(scheme_type=batch.scheme_type)
        target_obj = None
        match_type = 'none'
        if tid_val:
            target_obj = qs.filter(tid_number=tid_val).first()
            if target_obj:
                match_type = 'tid'
        if not target_obj and policy_val and name_val:
            target_obj = qs.filter(policy_no=policy_val).filter(
                Q(patient__full_name__iexact=name_val) | Q(patient_name_raw__iexact=name_val)
            ).first()
            if target_obj:
                match_type = 'policy_name'

        status_val = _tpa_normalize_status(cell(row_dict, 'status'))
        # Use the pre-summed total for this TID (covers files with one row
        # per claim line item) rather than just this row's own amount.
        amount_val = tid_totals[tid_val] if tid_val else _tpa_parse_amount(cell(row_dict, 'approved_amount'))
        package_val = cell(row_dict, 'package_code')
        package_val = str(package_val).strip() if package_val not in (None, '') else None
        admission_val = _tpa_parse_date(cell(row_dict, 'admission_date'))
        scheme_name_val = cell(row_dict, 'scheme_name')
        scheme_name_val = str(scheme_name_val).strip() if scheme_name_val not in (None, '') else None
        scheme_obj = None
        if scheme_name_val:
            scheme_obj = TPAScheme.objects.filter(scheme_type=batch.scheme_type, name__iexact=scheme_name_val).first()

        if target_obj:
            changes = {}

            def apply(field, newval):
                if newval is None:
                    return
                oldval = getattr(target_obj, field)
                if str(oldval or '') != str(newval):
                    changes[field] = {'before': _tpa_json_safe(oldval), 'after': _tpa_json_safe(newval)}
                    setattr(target_obj, field, newval)

            apply('status', status_val)
            apply('approved_amount', amount_val)
            apply('package_code', package_val)
            apply('admission_date', admission_val)
            if tid_val and tid_val != target_obj.tid_number:
                apply('tid_number', tid_val)
            if policy_val and policy_val != target_obj.policy_no:
                apply('policy_no', policy_val)
            if scheme_name_val:
                apply('scheme_name_raw', scheme_name_val)
            if scheme_obj and target_obj.scheme_id != scheme_obj.id:
                changes['scheme'] = {'before': target_obj.scheme_id, 'after': scheme_obj.id}
                target_obj.scheme = scheme_obj

            if changes:
                target_obj.save()

            TPAImportRow.objects.create(
                batch=batch, row_number=row_num, raw_data=raw_data,
                match_type=match_type, tpa_patient=target_obj,
                action='updated', field_changes=changes,
            )
            updated += 1
            if tid_val:
                tid_seen[tid_val] = target_obj
        else:
            new_obj = TPAPatient.objects.create(
                scheme_type=batch.scheme_type,
                patient_name_raw=name_val,
                tid_number=tid_val,
                policy_no=policy_val,
                package_code=package_val or '',
                approved_amount=amount_val,
                status=status_val or 'pending',
                admission_date=admission_val,
                scheme_name_raw=scheme_name_val or '',
                scheme=scheme_obj,
                is_draft=True,
            )
            TPAImportRow.objects.create(
                batch=batch, row_number=row_num, raw_data=raw_data,
                match_type='none', tpa_patient=new_obj, action='created',
            )
            created += 1
            if tid_val:
                tid_seen[tid_val] = new_obj

    wb.close()
    batch.updated_count = updated
    batch.created_count = created
    batch.merged_count = merged
    batch.skipped_count = skipped
    batch.status = 'done'
    batch.save(update_fields=['updated_count', 'created_count', 'merged_count', 'skipped_count', 'status'])


_TPA_UNDO_CASTERS = {
    'approved_amount': lambda v: Decimal(str(v)) if v not in (None, '') else None,
    'admission_date': lambda v: _dt.date.fromisoformat(v) if v else None,
}


def _undo_tpa_import(batch):
    for row in batch.rows.select_related('tpa_patient').all():
        if row.action == 'updated' and row.tpa_patient_id and row.field_changes:
            tp = row.tpa_patient
            changed = False
            for field, diff in row.field_changes.items():
                before = diff.get('before')
                try:
                    if field == 'scheme':
                        tp.scheme_id = before
                    else:
                        caster = _TPA_UNDO_CASTERS.get(field, lambda v: v)
                        setattr(tp, field, caster(before))
                    changed = True
                except Exception:
                    continue
            if changed:
                tp.save()
        elif row.action == 'created' and row.tpa_patient_id:
            tp = row.tpa_patient
            if tp.is_draft:
                tp.delete()


@login_required
def tpa_import_upload(request):
    if request.method == 'POST':
        f = request.FILES.get('excel_file')
        scheme_type = request.POST.get('scheme_type', 'government')
        if not f:
            messages.error(request, 'Choose an Excel file to import.')
        else:
            batch = TPAImportBatch.objects.create(
                file=f, original_filename=f.name, scheme_type=scheme_type,
                imported_by=request.user, status='mapping',
            )
            return redirect('hms:tpa_import_map', batch_id=batch.id)

    recent_batches = TPAImportBatch.objects.select_related('imported_by')[:10]
    return render(request, 'tpa/import_upload.html', {
        'scheme_type_choices': TPA_SCHEME_TYPE_CHOICES,
        'recent_batches': recent_batches,
    })


@login_required
def tpa_import_map(request, batch_id):
    batch = get_object_or_404(TPAImportBatch, id=batch_id)
    try:
        headers = _tpa_read_headers(batch.file.path)
    except Exception as e:
        messages.error(request, f'Could not read that Excel file: {e}')
        return redirect('hms:tpa_import_upload')

    if request.method == 'POST':
        mapping = {}
        for target, _label in TPA_IMPORT_TARGET_FIELDS:
            val = request.POST.get(f'map_{target}', '')
            if val:
                mapping[target] = val

        if 'tid' not in mapping and 'policy_no' not in mapping:
            messages.error(request, 'Map at least TID or Policy No./Beneficiary ID so rows can be matched.')
        else:
            batch.mapping = mapping
            batch.save(update_fields=['mapping'])
            TPAImportMapping.objects.update_or_create(
                name=f'default_{batch.scheme_type}', defaults={'mapping': mapping}
            )
            try:
                _run_tpa_import(batch)
            except Exception as e:
                batch.status = 'failed'
                batch.error_message = f'Import stopped due to an error: {e}'
                batch.save(update_fields=['status', 'error_message'])
            return redirect('hms:tpa_import_summary', batch_id=batch.id)

    preset = TPAImportMapping.objects.filter(name=f'default_{batch.scheme_type}').first()
    saved_mapping = {k: v for k, v in preset.mapping.items() if v in headers} if preset else {}
    guessed = _tpa_guess_mapping(headers)
    guessed.update(saved_mapping)

    return render(request, 'tpa/import_map.html', {
        'batch': batch,
        'headers': headers,
        'target_fields': TPA_IMPORT_TARGET_FIELDS,
        'guessed': guessed,
    })


@login_required
def tpa_import_summary(request, batch_id):
    batch = get_object_or_404(TPAImportBatch, id=batch_id)
    rows = batch.rows.select_related('tpa_patient').all()
    return render(request, 'tpa/import_summary.html', {'batch': batch, 'rows': rows})


@login_required
def tpa_import_batch_list(request):
    batches = TPAImportBatch.objects.select_related('imported_by', 'reverted_by').all()
    return render(request, 'tpa/import_batch_list.html', {'batches': batches})


@login_required
def tpa_import_undo(request, batch_id):
    batch = get_object_or_404(TPAImportBatch, id=batch_id)
    if request.method == 'POST':
        if batch.status != 'done':
            messages.error(request, 'Only a completed import can be reverted.')
        else:
            _undo_tpa_import(batch)
            batch.status = 'reverted'
            batch.reverted_by = request.user
            batch.reverted_at = timezone.now()
            batch.save(update_fields=['status', 'reverted_by', 'reverted_at'])
            messages.success(request, f'Import #{batch.id} reverted — changes rolled back.')
    return redirect('hms:tpa_import_summary', batch_id=batch.id)

