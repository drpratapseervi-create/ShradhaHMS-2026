import os
import json

import logging

from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse, HttpResponse
from django.db.models import Q
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_POST

from ..decorators import role_required
from ..models import Patient, Doctor, Consultation, MedicalImage, USGReport
from ..forms import USGReportForm
from ..utils import render_to_pdf
from ..templatetags.hms_extras import comma_split, usg_findings_line_parts
from ..services.whatsapp import send_usg_report_pdf, WhatsAppSendError

logger = logging.getLogger("hms.views.imaging")


@login_required
def upload_medical_image(request, patient_id):
    patient = get_object_or_404(Patient, id=patient_id)
    if request.method == "POST":
        image = request.FILES.get("image")
        if image:
            MedicalImage.objects.create(
                patient=patient,
                image_type=request.POST.get("image_type"),
                title=request.POST.get("title"),
                image=image,
                report_text=request.POST.get("report_text"),
                uploaded_by=request.user,
            )
            messages.success(request, "Medical image uploaded successfully.")
        return redirect("hms:dashboard")
    return render(request, "medical/upload_image.html", {"patient": patient})


@login_required
def upload_medical_image_consultation(request, patient_id, consultation_id):
    patient      = get_object_or_404(Patient, id=patient_id)
    consultation = get_object_or_404(Consultation, id=consultation_id)

    if request.method == "POST":
        image_type  = request.POST.get("image_type", "XRAY")
        title       = request.POST.get("title", "").strip()
        image       = request.FILES.get("image")
        report_text = request.POST.get("report_text", "").strip()

        if image and title:
            MedicalImage.objects.create(
                patient      = patient,
                consultation = consultation,
                image_type   = image_type,
                title        = title,
                image        = image,
                report_text  = report_text,
                uploaded_by  = request.user,
            )
            messages.success(request, f"Image '{title}' uploaded successfully.")
        else:
            messages.error(request, "Title and image file are required.")

    return redirect(
        "hms:start_consultation",
        appointment_id=consultation.appointment.id,
    )


@login_required
def delete_medical_image(request, image_id):
    img = get_object_or_404(MedicalImage, id=image_id)

    appointment_id = request.POST.get("appointment_id")
    redirect_to    = request.POST.get("redirect_to", "radiology")

    if img.image:
        if os.path.isfile(img.image.path):
            os.remove(img.image.path)

    img.delete()
    messages.success(request, "Image deleted successfully.")

    if redirect_to == "consultation" and appointment_id:
        return redirect("hms:start_consultation", appointment_id=appointment_id)

    return redirect("hms:radiology_upload")


@login_required
@role_required("doctor", "admin", "nursing", "laboratory")
def radiology_upload(request):
    from hms.dicom import upload_dicom

    search_query = request.GET.get("q", "").strip()
    if search_query:
        patients = Patient.objects.filter(
            Q(full_name__icontains=search_query) |
            Q(uhid__icontains=search_query) |
            Q(mobile_no__icontains=search_query)
        ).order_by("-id")[:50]
    else:
        patients = Patient.objects.all().order_by("-id")[:50]

    if request.method == "POST":
        patient_id  = request.POST.get("patient_id")
        image_type  = request.POST.get("image_type")
        title       = request.POST.get("title", "").strip()
        image       = request.FILES.get("image")
        report_text = request.POST.get("report_text", "").strip()

        errors = []
        if not patient_id: errors.append("Please select a patient.")
        if not title:       errors.append("Please enter a title.")
        if not image:       errors.append("Please select an image file.")

        if errors:
            for e in errors:
                messages.error(request, e)
        else:
            patient = get_object_or_404(Patient, id=patient_id)

            # ✅ DICOM file (.dcm) — upload to Orthanc
            if image.name.lower().endswith('.dcm'):
                dicom_bytes = image.read()
                instance_id = upload_dicom(dicom_bytes)
                if instance_id:
                    MedicalImage.objects.create(
                        patient           = patient,
                        image_type        = image_type,
                        title             = title,
                        report_text       = report_text,
                        uploaded_by       = request.user,
                        dicom_instance_id = instance_id,
                        is_dicom          = True,
                    )
                    messages.success(request,
                        f"✅ DICOM uploaded to Orthanc for "
                        f"{patient.full_name} ({patient.uhid}) "
                        f"— ID: {instance_id[:8]}..."
                    )
                else:
                    messages.error(request,
                        "DICOM upload to Orthanc failed. "
                        "Make sure Orthanc is running on port 8042."
                    )

            # Regular image (JPG/PNG/GIF/WEBP)
            else:
                allowed_types = ["image/jpeg", "image/png",
                                 "image/gif",  "image/webp"]
                if image.content_type not in allowed_types:
                    messages.error(request,
                        "Only image files (JPG, PNG, GIF, WEBP) "
                        "or DICOM (.dcm) files are allowed."
                    )
                else:
                    MedicalImage.objects.create(
                        patient     = patient,
                        image_type  = image_type,
                        title       = title,
                        image       = image,
                        report_text = report_text,
                        uploaded_by = request.user,
                        is_dicom    = False,
                    )
                    messages.success(request,
                        f"✅ Image uploaded for "
                        f"{patient.full_name} ({patient.uhid})"
                    )

            return redirect("hms:radiology_upload")

    return render(request, "lab/radiology_upload.html", {
        "patients":       patients,
        "search_query":   search_query,
        "recent_uploads": MedicalImage.objects.select_related("patient")
                          .order_by("-created_at")[:15],
    })


@login_required
def usg_report_list(request):
    """List USG reports — with patient search filter."""
    q = request.GET.get("q", "").strip()
    reports = USGReport.objects.select_related("patient", "reporting_doctor")

    if q:
        reports = reports.filter(
            Q(patient__full_name__icontains=q) |
            Q(patient__uhid__icontains=q)      |
            Q(report_no__icontains=q)
        )

    return render(request, "hms/usg/usg_report_list.html", {
        "reports": reports[:100],
        "q":       q,
    })


USG_ADVICE_QUICK_OPTIONS = [
    "LFT", "S. Amylase", "X-ray KUB", "IVP", "CT Urography",
    "CECT Abdomen & Pelvis", "MRCP", "TVS", "UPT", "Urine R/E",
]


def _usg_impression_groups():
    from ..models import USGImpressionOption
    options = USGImpressionOption.objects.filter(is_active=True).order_by("category", "sort_order", "text")
    groups = []
    current_key = None
    for opt in options:
        label = opt.get_category_display()
        if label != current_key:
            groups.append((label, []))
            current_key = label
        groups[-1][1].append(opt.text)
    return groups


@login_required
def usg_report_create(request, patient_id=None, bill_item_id=None):
    """Create a new USG report — optionally pre-linked to patient / bill item."""
    patient   = None
    bill_item = None

    if patient_id:
        patient = get_object_or_404(Patient, pk=patient_id)
    if bill_item_id:
        from ..models import InvestigationBillItem
        bill_item = get_object_or_404(InvestigationBillItem, pk=bill_item_id)
        if not patient:
            patient = bill_item.bill.patient

    initial = {}
    if patient:
        initial["patient"] = patient
    if bill_item:
        initial["bill_item"] = bill_item
    initial["findings_text"] = USGReport.default_findings_text(
        "ABDOMEN_PELVIS", gender=patient.gender if patient else None
    )
    initial["clinical_indication"] = "Pain abdomen"
    default_doctor = Doctor.objects.filter(full_name__icontains="Pratap Senecha").first()
    if default_doctor:
        initial.setdefault("referred_by", default_doctor)
        initial.setdefault("reporting_doctor", default_doctor)

    is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    form = USGReportForm(request.POST or None, initial=initial)

    if request.method == "POST" and form.is_valid():
        report = form.save(commit=False)
        report.created_by = request.user
        report.save()
        if is_ajax:
            return JsonResponse({"success": True, "report_id": report.pk})
        messages.success(request, f"USG Report {report.report_no} saved successfully.")
        return redirect("hms:usg_report_print", pk=report.pk)
    elif request.method == "POST" and is_ajax:
        return JsonResponse({"success": False, "error": form.errors.as_text()}, status=400)

    return render(request, "hms/usg/usg_report_form.html", {
        "form":               form,
        "patient":            patient,
        "bill_item":          bill_item,
        "title":              "New USG Report",
        "findings_templates": json.dumps(USGReport.FINDINGS_TEMPLATES),
        "measurement_fields": USGReport.MEASUREMENT_FIELDS,
        "pelvic_organs_texts": json.dumps(USGReport.pelvic_organs_texts()),
        "advice_quick_options": USG_ADVICE_QUICK_OPTIONS,
        "impression_groups": _usg_impression_groups(),
    })


@login_required
def usg_report_edit(request, pk):
    report = get_object_or_404(USGReport, pk=pk)
    form   = USGReportForm(request.POST or None, instance=report)
    is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"

    if request.method == "POST" and form.is_valid():
        form.save()
        if is_ajax:
            return JsonResponse({"success": True, "report_id": report.pk})
        messages.success(request, f"Report {report.report_no} updated.")
        return redirect("hms:usg_report_print", pk=report.pk)
    elif request.method == "POST" and is_ajax:
        return JsonResponse({"success": False, "error": form.errors.as_text()}, status=400)

    return render(request, "hms/usg/usg_report_form.html", {
        "form":               form,
        "report":             report,
        "title":              f"Edit {report.report_no}",
        "findings_templates": json.dumps(USGReport.FINDINGS_TEMPLATES),
        "measurement_fields": USGReport.MEASUREMENT_FIELDS,
        "pelvic_organs_texts": json.dumps(USGReport.pelvic_organs_texts()),
        "advice_quick_options": USG_ADVICE_QUICK_OPTIONS,
        "impression_groups": _usg_impression_groups(),
    })


@login_required
def usg_report_print(request, pk):
    """Render a print-ready USG report page."""
    report = get_object_or_404(
        USGReport.objects.select_related(
            "patient", "reporting_doctor", "referred_by", "consultation"
        ),
        pk=pk
    )
    return render(request, "hms/usg/usg_report_print.html", {
        "report": report,
    })


@login_required
def usg_report_pdf(request, pk):
    """Generate PDF of USG report via xhtml2pdf."""
    report = get_object_or_404(
        USGReport.objects.select_related(
            "patient", "reporting_doctor", "referred_by"
        ),
        pk=pk
    )
    return render_to_pdf("hms/usg/usg_report_print.html", {"report": report})


@login_required
@require_POST
def usg_report_send_whatsapp(request, pk):
    report = get_object_or_404(
        USGReport.objects.select_related(
            "patient", "reporting_doctor", "referred_by"
        ),
        pk=pk
    )
    pdf_bytes = render_to_pdf("hms/usg/usg_report_print.html", {"report": report}).content

    try:
        send_usg_report_pdf(report, pdf_bytes)
    except ValueError as exc:
        return JsonResponse({"success": False, "error": str(exc)}, status=400)
    except WhatsAppSendError:
        logger.exception("Failed to send USG report PDF via WhatsApp for report %s", report.id)
        return JsonResponse({"success": False, "error": "Failed to send via WhatsApp. Please try again."}, status=502)

    return JsonResponse({"success": True})


def _shade_cell(cell, color_hex):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def _build_usg_report_docx(report):
    """Word version of the USG print report — same content/emphasis rules
    (organ labels bold, Impression numbered + bold, Advice bold, 5cm top
    space reserved for letterhead) as hms/usg/usg_report_print.html."""
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()
    section = doc.sections[0]
    # Printed on pre-printed hospital letterhead — same 5cm reserved space
    # as the print template's @page margin-top.
    section.top_margin = Cm(5)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(1.5)

    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)

    # ── Patient Details table ──
    rows = [
        ("Patient Name", (report.patient.full_name or "").title(), "Age / Sex",
         f"{report.patient.age_years} Yrs" if report.patient.age_years else "-",
         f" / {report.patient.gender}"),
        ("UHID", report.patient.uhid, "Referred By",
         report.referred_by.full_name if report.referred_by else "-", ""),
        ("Scan Type", report.get_scan_type_display(), "Clinical Indication",
         report.clinical_indication or "-", ""),
        ("Report No.", report.report_no, "Date",
         report.report_date.strftime("%d-%m-%Y") if report.report_date else "-", ""),
    ]
    table = doc.add_table(rows=len(rows), cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, (lbl1, val1, lbl2, val2, suffix2) in enumerate(rows):
        cells = table.rows[r].cells
        cells[0].paragraphs[0].add_run(lbl1).bold = True
        cells[1].paragraphs[0].add_run(str(val1 or "-"))
        cells[2].paragraphs[0].add_run(lbl2).bold = True
        cells[3].paragraphs[0].add_run(str(val2 or "-") + suffix2)
        _shade_cell(cells[0], "F5F5F5")
        _shade_cell(cells[2], "F5F5F5")

    doc.add_paragraph()

    # ── Title bar ──
    title_table = doc.add_table(rows=1, cols=1)
    title_cell = title_table.rows[0].cells[0]
    _shade_cell(title_cell, "0369A1")
    p = title_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"ULTRASONOGRAPHY {report.get_scan_type_display().upper()}")
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    doc.add_paragraph()

    # ── Findings — organ labels + manually-added lines bold, standard
    #    narrative normal, same classification as the print template ──
    for line in (report.findings_text or "-").replace("\r\n", "\n").split("\n"):
        para = doc.add_paragraph()
        for bold, text in usg_findings_line_parts(line):
            para.add_run(text).bold = bold

    doc.add_paragraph()

    # ── Impression — numbered vertical list, always bold ──
    impression_items = comma_split(report.impression)
    heading = doc.add_paragraph()
    heading.add_run("IMPRESSION").bold = True
    if not impression_items:
        doc.add_paragraph().add_run("-").bold = True
    else:
        for i, item in enumerate(impression_items, start=1):
            doc.add_paragraph().add_run(f"{i}. {item}").bold = True

    # ── Advice — bold, same as print ──
    if report.advice:
        doc.add_paragraph()
        advice_p = doc.add_paragraph()
        advice_p.add_run("ADVICE: ").bold = True
        advice_p.add_run(report.advice).bold = True

    # ── Signature ──
    doc.add_paragraph()
    doc.add_paragraph()
    sig = doc.add_paragraph()
    sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doctor_name = f"Dr. {report.reporting_doctor.full_name if report.reporting_doctor else 'Pratap Senecha'}"
    qualification = (
        report.reporting_doctor.qualification if report.reporting_doctor and report.reporting_doctor.qualification
        else ("MBBS, MS" if not report.reporting_doctor else "")
    )
    sig.add_run(doctor_name + (f", {qualification}" if qualification else "")).bold = True
    reg = doc.add_paragraph()
    reg.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    reg.add_run(report.reporting_doctor.registration_no if report.reporting_doctor else "RMC No-27994")

    return doc


@login_required
def usg_report_download_word(request, pk):
    """Download the USG report as an editable .docx (same content/emphasis
    as the print/PDF version)."""
    report = get_object_or_404(
        USGReport.objects.select_related("patient", "reporting_doctor", "referred_by"),
        pk=pk
    )
    doc = _build_usg_report_docx(report)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    response["Content-Disposition"] = f'attachment; filename="{report.report_no or ("USG-" + str(report.pk))}.docx"'
    doc.save(response)
    return response


@login_required
def usg_report_delete(request, pk):
    report = get_object_or_404(USGReport, pk=pk)
    if request.method == "POST":
        report.delete()
        messages.success(request, "USG Report deleted.")
        return redirect("hms:usg_report_list")
    return render(request, "hms/usg/usg_report_confirm_delete.html", {"report": report})

