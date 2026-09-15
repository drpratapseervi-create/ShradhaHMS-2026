import csv
from datetime import datetime

import openpyxl
from reportlab.platypus import SimpleDocTemplate, Table
from reportlab.lib import colors

from django.shortcuts import render
from django.http import HttpResponse
from django.db.models import Sum
from django.contrib.auth.decorators import login_required
from django.utils import timezone

from ..decorators import role_required
from ..models import (
    ProcedureBill, Consultation, InvestigationBill, InvestigationBillItem,
    IPDAdvance, IPDAdmission, DischargeBill, Appointment, Expense, Doctor,
)
from ._shared import logger


def daily_report(request):
    today         = timezone.now().date()
    selected_date = request.GET.get('date')

    if selected_date:
        try:
            filter_date = datetime.strptime(selected_date, '%Y-%m-%d').date()
        except ValueError:
            filter_date = today
    else:
        filter_date = today

    # ---------------- LAB COLLECTION ----------------
    # FIX: added bill__paid=True so only collected (paid) bills count
    def inv_by_dept(dept_code):
        return InvestigationBillItem.objects.filter(
            bill__created_at__date=filter_date,
            bill__paid=True,
            investigation__category__dept_code__iexact=dept_code
        ).aggregate(total=Sum('price'))['total'] or 0

    radiology  = inv_by_dept('RADIOLOGY')
    ecg        = inv_by_dept('ECG')
    cardiology = inv_by_dept('CARDIOLOGY')
    histo      = inv_by_dept('HISTOPATHOLOGY')
    biochem    = inv_by_dept('BIOCHEMISTRY')
    hematology = inv_by_dept('HEMATOLOGY')
    microbio   = inv_by_dept('MICROBIOLOGY')
    endoscopy  = inv_by_dept('ENDOSCOPY')

    lab = radiology + ecg + cardiology + histo + biochem + hematology + microbio + endoscopy

    # ---------------- OPD COLLECTION ----------------
    # FIX: is_paid is now correctly set in appointment_create
    opd = Appointment.objects.filter(
        date=filter_date,
        is_paid=True
    ).aggregate(total=Sum('fee'))['total'] or 0

    # ---------------- PROCEDURE COLLECTION ----------------
    procedure = ProcedureBill.objects.filter(
        created_at__date=filter_date
    ).aggregate(total=Sum('net_amount'))['total'] or 0

    # ---------------- IPD ADVANCE ----------------
    advance = IPDAdvance.objects.filter(
        date__date=filter_date
    ).aggregate(total=Sum('amount'))['total'] or 0

    # ---------------- DISCHARGE COLLECTION ----------------
    discharge = DischargeBill.objects.filter(
        paid_at__date=filter_date,
        is_paid=True
    ).aggregate(total=Sum('final_amount'))['total'] or 0

    # ---------------- EXPENSES ----------------
    expenses = Expense.objects.filter(
        date=filter_date
    ).aggregate(total=Sum('amount'))['total'] or 0

    # ---------------- FINAL CALCULATION ----------------
    total_collection = opd + lab + procedure + advance + discharge
    net_collection   = total_collection - expenses

    return render(request, "daily_report.html", {
        "selected_date":    filter_date,
        "opd":              opd,
        "radiology":        radiology,
        "ecg":              ecg,
        "cardiology":       cardiology,
        "histo":            histo,
        "biochem":          biochem,
        "hematology":       hematology,
        "microbio":         microbio,
        "endoscopy":        endoscopy,
        "lab":              lab,
        "procedure":        procedure,
        "advance":          advance,
        "discharge":        discharge,
        "expenses":         expenses,
        "total_collection": total_collection,
        "net_collection":   net_collection,
    })

    from django.shortcuts import get_object_or_404


def get_dept_display(dept_code):
    mapping = {
        "RADIOLOGY": "Radiology",
        "HISTOPATHOLOGY": "Histopathology",
        "BIOCHEMISTRY": "Biochemistry",
        "HEMATOLOGY": "Hematology",
        "MICROBIOLOGY": "Microbiology",
        "CARDIOLOGY": "Cardiology",
        "ECG": "ECG",
        "ENDOSCOPY": "Endoscopy",
        "OTHER": "Other",
    }
    return mapping.get(dept_code, "Other")


def get_all_payment_data(request):
    
    from_date = request.GET.get("from_date")
    to_date = request.GET.get("to_date")
    doctor = request.GET.get("doctor")
    dept = request.GET.get("department")

    data = []

    # ================= PROCEDURE =================
    procedures = ProcedureBill.objects.select_related(
        "patient", "consultant", "department"
    )

    if from_date and to_date:
        procedures = procedures.filter(created_at__date__range=[from_date, to_date])

    if doctor:
        procedures = procedures.filter(consultant_id=doctor)

    if dept:
        procedures = procedures.filter(department_id=dept)

    for p in procedures:
        local_p = timezone.localtime(p.created_at)
        data.append({
            "receipt_no": f"PB-{p.id}",
            "date": local_p,
            "time": local_p.strftime("%I:%M %p"),
            "patient": p.patient.full_name,
            "type": "Procedure",
            "doctor": p.consultant.full_name if p.consultant else "",
            "department": p.department.name if p.department else "",
            "amount": float(p.net_amount)
        })

    # ================= OPD =================
    consultations = Consultation.objects.select_related(
        "appointment__patient", "appointment__doctor"
    )

    if from_date and to_date:
        consultations = consultations.filter(created_at__date__range=[from_date, to_date])

    for c in consultations:
        local_created = timezone.localtime(c.created_at)
        data.append({
            "receipt_no": f"OPD-{c.id}",
            "date": local_created,
            "time": local_created.strftime("%I:%M %p"),
            "patient": c.appointment.patient.full_name if c.appointment and c.appointment.patient else "",
            "type": "OPD",
            "doctor": c.appointment.doctor.full_name if c.appointment and c.appointment.doctor else "",
            "department": "OPD",
            "amount": 100
        })

    # ================= LAB (InvestigationBill) =================
    try:
        lab_bills = InvestigationBill.objects.prefetch_related(
            "items__investigation__category"
        ).select_related("patient", "consultation__appointment__doctor")

        if from_date and to_date:
            lab_bills = lab_bills.filter(created_at__date__range=[from_date, to_date])

        for l in lab_bills:
            dept_set = set()
            for item in l.items.all():
                if item.investigation and item.investigation.category:
                    dept_set.add(item.investigation.category.dept_code)
            dept_names = [get_dept_display(code) for code in dept_set]
            dept_display = ", ".join(dept_names) if dept_names else "Laboratory"

            amount = float(l.net_amount) if l.net_amount else float(l.total_amount)
            local_l = timezone.localtime(l.created_at)

            lab_doctor = ""
            if l.consultation and l.consultation.appointment and l.consultation.appointment.doctor:
                lab_doctor = l.consultation.appointment.doctor.full_name
            elif l.patient:
                last_appt = Appointment.objects.filter(patient=l.patient).order_by("-id").first()
                if last_appt and last_appt.doctor:
                    lab_doctor = last_appt.doctor.full_name

            data.append({
                "receipt_no": f"LAB-{l.id}",
                "date": local_l,
                "time": local_l.strftime("%I:%M %p"),
                "patient": l.patient.full_name if l.patient else "",
                "type": "Lab",
                "doctor": lab_doctor,
                "department": dept_display,
                "amount": amount
            })
    except Exception as e:
        logger.error(f"Lab billing fetch error: {e}")

    # ================= IPD ADVANCE =================
    try:
        advances = IPDAdvance.objects.select_related("patient")
        if from_date and to_date:
            advances = advances.filter(date__date__range=[from_date, to_date])
        for a in advances:
            local_date = timezone.localtime(a.date)
            adm = IPDAdmission.objects.filter(patient=a.patient).order_by("-id").first()
            adv_doctor = adm.doctor.full_name if adm and adm.doctor else ""
            adv_dept   = adm.department.name  if adm and adm.department else "IPD"
            data.append({
                "receipt_no": f"ADV-{a.pk:05d}",
                "date": local_date,
                "time": local_date.strftime("%I:%M %p"),
                "patient": a.patient.full_name if a.patient else "",
                "type": "Advance",
                "doctor": adv_doctor,
                "department": adv_dept,
                "amount": float(a.amount)
            })
    except Exception as e:
        logger.error(f"IPD advance fetch error: {e}")

    # ================= DISCHARGE FINAL PAYMENT (FPR) =================
    try:
        discharge_bills = DischargeBill.objects.select_related("patient").filter(is_paid=True)
        if from_date and to_date:
            discharge_bills = discharge_bills.filter(paid_at__date__range=[from_date, to_date])
        for d in discharge_bills:
            paid_at = d.paid_at or d.created_at
            local_paid = timezone.localtime(paid_at)
            adm = IPDAdmission.objects.filter(patient=d.patient).order_by("-id").first()
            dis_doctor = adm.doctor.full_name if adm and adm.doctor else ""
            dis_dept   = adm.department.name  if adm and adm.department else "IPD"
            data.append({
                "receipt_no": f"FPR-{d.id:05d}",
                "date": local_paid,
                "time": local_paid.strftime("%I:%M %p"),
                "patient": d.patient.full_name if d.patient else "",
                "type": "Discharge",
                "doctor": dis_doctor,
                "department": dis_dept,
                "amount": float(d.final_amount)
            })
    except Exception as e:
        logger.error(f"Discharge final payment fetch error: {e}")

    # SORT
    data = sorted(data, key=lambda x: x["date"], reverse=True)

    return data


@login_required
@role_required("admin")
def all_reports(request):
    data = get_all_payment_data(request)

    type_filter = request.GET.get("type_filter", "")
    if type_filter:
        data = [r for r in data if r["type"].lower() == type_filter.lower()]

    dept_filter = request.GET.get("dept_filter", "")
    if dept_filter:
        df = dept_filter.lower()
        item_level_filters = {
            "xray":     ["x-ray", "x ray", "xray", "chest x", "spine", "kub", "skull"],
            "usg":      ["usg", "ultrasound", "sonography", "usg abdomen", "usg pelvis"],
            "serology": ["hiv", "hbsag", "hcv", "vdrl", "widal", "aso", "crp", "typhoid", "tridot"],
        }
        if df in item_level_filters:
            keywords = item_level_filters[df]
            matching_ids = set()
            for row in data:
                if row["type"] == "Lab":
                    try:
                        bill_id = int(row["receipt_no"].replace("LAB-", ""))
                        items = InvestigationBillItem.objects.filter(
                            bill_id=bill_id
                        ).select_related("investigation")
                        for item in items:
                            name = item.investigation.name.lower() if item.investigation else ""
                            if any(k in name for k in keywords):
                                matching_ids.add(row["receipt_no"])
                                break
                    except Exception:
                        pass
            data = [r for r in data if r["receipt_no"] in matching_ids]
        else:
            data = [r for r in data if df in r["department"].lower()]

    doctor_filter = request.GET.get("doctor_filter", "")
    if doctor_filter:
        data = [r for r in data if doctor_filter.lower() in r["doctor"].lower()]

    total   = sum(x["amount"] for x in data)
    doctors = Doctor.objects.order_by("full_name")

    return render(request, "reports/all_reports.html", {
        "data":          data,
        "total":         total,
        "doctors":       doctors,
        "type_filter":   type_filter,
        "dept_filter":   dept_filter,
        "doctor_filter": doctor_filter,
        "from_date":     request.GET.get("from_date", ""),
        "to_date":       request.GET.get("to_date", ""),
    })


@login_required
@role_required("admin")
def export_all_csv(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="all_payments.csv"'

    writer = csv.writer(response)
    writer.writerow(['Receipt No', 'Date', 'Time', 'Patient', 'Type', 'Doctor', 'Department', 'Amount'])

    data = get_all_payment_data(request)

    for row in data:
        writer.writerow([
            row["receipt_no"],
            row["date"].strftime("%d-%m-%Y"),
            row["time"],
            row["patient"],
            row["type"],
            row["doctor"],
            row["department"],
            row["amount"]
        ])

    return response


def export_all_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All Payments"

    from openpyxl.styles import Font, PatternFill, Alignment
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0D3B66")

    headers = ['Receipt No', 'Date', 'Time', 'Patient', 'Type', 'Doctor', 'Department', 'Amount (Rs)']
    ws.append(headers)
    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    data = get_all_payment_data(request)
    for row in data:
        ws.append([
            row["receipt_no"],
            row["date"].strftime("%d-%m-%Y"),
            row["time"],
            row["patient"],
            row["type"],
            row["doctor"],
            row["department"],
            row["amount"]
        ])

    col_widths = [14, 13, 10, 20, 12, 22, 35, 14]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=all_payments.xlsx'
    wb.save(response)
    return response


def export_all_pdf(request):
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.units import mm
    from reportlab.platypus import TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import ParagraphStyle

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="all_payments.pdf"'

    doc = SimpleDocTemplate(
        response, pagesize=landscape(A4),
        leftMargin=10*mm, rightMargin=10*mm,
        topMargin=12*mm, bottomMargin=12*mm
    )

    title_style = ParagraphStyle('t', fontSize=14, fontName='Helvetica-Bold', alignment=1, spaceAfter=6)
    sub_style   = ParagraphStyle('s', fontSize=9,  fontName='Helvetica',      alignment=1, spaceAfter=10)

    rows = get_all_payment_data(request)
    total = sum(r["amount"] for r in rows)
    from_date = request.GET.get("from_date", "")
    to_date   = request.GET.get("to_date", "")
    period = f"{from_date} to {to_date}" if from_date and to_date else "All Dates"

    table_data = [['Receipt No', 'Date', 'Time', 'Patient', 'Type', 'Doctor', 'Department', 'Amount (Rs)']]
    for r in rows:
        table_data.append([
            r["receipt_no"], r["date"].strftime("%d-%m-%Y"), r["time"],
            r["patient"], r["type"], r["doctor"], r["department"],
            f"{r['amount']:,.2f}"
        ])
    table_data.append(['', '', '', '', '', '', 'TOTAL', f"{total:,.2f}"])

    col_widths = [22*mm, 22*mm, 16*mm, 38*mm, 20*mm, 40*mm, 60*mm, 26*mm]
    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND',    (0,0),  (-1,0),  colors.HexColor('#0D3B66')),
        ('TEXTCOLOR',     (0,0),  (-1,0),  colors.white),
        ('FONTNAME',      (0,0),  (-1,0),  'Helvetica-Bold'),
        ('FONTSIZE',      (0,0),  (-1,0),  8),
        ('ALIGN',         (0,0),  (-1,0),  'CENTER'),
        ('FONTSIZE',      (0,1),  (-1,-2), 7.5),
        ('ROWBACKGROUNDS',(0,1),  (-1,-2), [colors.white, colors.HexColor('#F0F4FF')]),
        ('BACKGROUND',    (0,-1), (-1,-1), colors.HexColor('#0D3B66')),
        ('TEXTCOLOR',     (0,-1), (-1,-1), colors.white),
        ('FONTNAME',      (0,-1), (-1,-1), 'Helvetica-Bold'),
        ('ALIGN',         (-1,0), (-1,-1), 'RIGHT'),
        ('GRID',          (0,0),  (-1,-1), 0.4, colors.HexColor('#CCCCCC')),
        ('VALIGN',        (0,0),  (-1,-1), 'MIDDLE'),
        ('TOPPADDING',    (0,0),  (-1,-1), 4),
        ('BOTTOMPADDING', (0,0),  (-1,-1), 4),
    ]))

    doc.build([
        Paragraph("Shradha Hospital & Multispeciality Centre", title_style),
        Paragraph(f"All Payment Receipts — {period}  |  Total: Rs.{total:,.2f}", sub_style),
        table
    ])
    return response

