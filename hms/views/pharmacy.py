"""Pharmacy counter: pending prescriptions, dispensing + billing, bills, H1 register, opening stock import."""

from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Q, Sum
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from ..decorators import role_required
from ..models import (
    Consultation, Doctor, DrugMaster, InventoryItem, IPDAdmission, Patient, PharmacyBill,
    PharmacyBillItem, PharmacyReturn, Prescription,
)
from ..pharmacy.dispensing import (
    cancel_bill, counter_collection, create_bill, ipd_medication_lines, ipd_pharmacy_total,
    item_stock_info, normalize_name, pending_consultations, prescription_lines, return_items,
)
from ..pharmacy.stock import InsufficientStock, receive_stock

PHARMACY_ROLES = ("pharmacy", "admin")
PENDING_LOOKBACK_DAYS = 3


def _parse_date(value, default):
    try:
        return datetime.strptime(value, "%Y-%m-%d").date() if value else default
    except ValueError:
        return default


@login_required
@role_required(*PHARMACY_ROLES)
def pharmacy_home(request):
    """Today's prescriptions waiting at the counter, and today's sales."""
    day = _parse_date(request.GET.get("date"), timezone.localdate())
    # Patients don't always reach the counter the same day, so look back a few days.
    pending, done = pending_consultations(day - timedelta(days=PENDING_LOOKBACK_DAYS - 1), day)
    bills = PharmacyBill.objects.filter(created_at__date=day).select_related("patient", "admission")
    collection, by_mode = counter_collection(day, day)
    admitted = (IPDAdmission.objects.filter(status="ADMITTED")
                .select_related("patient", "bed__ward", "doctor").order_by("-admission_date"))
    return render(request, "pharmacy/home.html", {
        "day": day,
        "is_today": day == timezone.localdate(),
        "pending": pending,
        "done": [r for r in done if r["c"].appointment.date == day],
        "lookback_days": PENDING_LOOKBACK_DAYS,
        "bills": bills,
        "collection": collection,
        "by_mode": by_mode,
        "admitted": [{"a": a, "pharmacy_total": ipd_pharmacy_total(a)} for a in admitted],
    })


def _line_items_from_post(post):
    """Read the counter table rows (only ticked rows with an item and quantity)."""
    ids, qtys = post.getlist("item_id"), post.getlist("quantity")
    rx_ids, include = post.getlist("prescription_id"), post.getlist("include")
    items = InventoryItem.objects.in_bulk([i for i in ids if i])
    rxs = Prescription.objects.in_bulk([r for r in rx_ids if r])
    lines = []
    for idx, (item_id, qty) in enumerate(zip(ids, qtys)):
        if str(idx) not in include or not item_id:
            continue
        try:
            qty = int(qty)
        except (TypeError, ValueError):
            continue
        rx_id = rx_ids[idx] if idx < len(rx_ids) else ""
        lines.append({
            "item": items.get(int(item_id)),
            "quantity": qty,
            "prescription": rxs.get(int(rx_id)) if rx_id else None,
        })
    return lines


@login_required
@role_required(*PHARMACY_ROLES)
def pharmacy_dispense(request, consultation_id=None, admission_id=None):
    """
    The counter screen, then bill + deduct stock. Prefilled from an OPD prescription,
    from an admission's medication orders (IPD issue, charged to the discharge bill),
    or blank for a walk-in.
    """
    consultation = admission = None
    if consultation_id:
        consultation = get_object_or_404(
            Consultation.objects.select_related("appointment__patient", "appointment__doctor"),
            id=consultation_id,
        )
    if admission_id:
        admission = get_object_or_404(
            IPDAdmission.objects.select_related("patient", "doctor", "bed__ward"), id=admission_id
        )
        if admission.status != "ADMITTED":
            messages.error(request, f"{admission.ipd_no} is discharged; medicines can't be charged to it.")
            return redirect("hms:pharmacy_home")

    if request.method == "POST":
        lines = _line_items_from_post(request.POST)
        try:
            discount = Decimal(request.POST.get("discount") or "0")
        except InvalidOperation:
            discount = Decimal("0")
        patient = consultation.appointment.patient if consultation else None
        if not patient and request.POST.get("patient_id"):
            patient = Patient.objects.filter(id=request.POST["patient_id"]).first()
        doctor = consultation.appointment.doctor if consultation else None
        if not doctor and request.POST.get("doctor_id"):
            doctor = Doctor.objects.filter(id=request.POST["doctor_id"]).first()
        try:
            bill = create_bill(
                lines=lines, user=request.user,
                payment_mode=request.POST.get("payment_mode", "CASH"),
                discount=discount, patient=patient, consultation=consultation,
                customer_name=request.POST.get("customer_name", "").strip(),
                customer_mobile=request.POST.get("customer_mobile", "").strip(),
                doctor=doctor, doctor_name=request.POST.get("doctor_name", "").strip(),
                admission=admission,
            )
        except InsufficientStock as e:
            messages.error(request, f"Not enough stock: {e}. Nothing was billed.")
        except ValueError as e:
            messages.error(request, str(e))
        else:
            if bill.is_ipd:
                messages.success(request, f"{bill.bill_no}: ₹{bill.net_amount} charged to {admission.ipd_no}'s discharge bill")
            else:
                messages.success(request, f"Bill {bill.bill_no} saved: ₹{bill.net_amount}")
            return redirect("hms:pharmacy_bill_print", bill_id=bill.id)

    if consultation:
        lines = prescription_lines(consultation)
    elif admission:
        lines = ipd_medication_lines(admission)
    else:
        lines = []
    h1_items = [l for l in lines if l["item"] and l["item"].schedule in ("H1", "X")]
    patient = consultation.appointment.patient if consultation else (admission.patient if admission else None)
    return render(request, "pharmacy/dispense.html", {
        "consultation": consultation,
        "admission": admission,
        "ipd_total": ipd_pharmacy_total(admission) if admission else None,
        "patient": patient,
        "lines": lines,
        "has_h1": bool(h1_items),
        "doctors": Doctor.objects.all().order_by("full_name"),
        "payment_modes": [m for m in PharmacyBill.PAYMENT_MODES if m[0] != "IPD"],
    })


@login_required
@role_required(*PHARMACY_ROLES)
def pharmacy_item_search(request):
    """JSON for the counter's 'add medicine' search (and substitutions)."""
    q = request.GET.get("q", "").strip()
    items = InventoryItem.objects.select_related("drug").order_by("name")
    if q:
        items = items.filter(Q(name__icontains=q) | Q(drug__name__icontains=q) | Q(drug__generic_name__icontains=q))
    results = []
    for item in items[:20]:
        info = item_stock_info(item)
        results.append({
            "id": item.id, "name": item.name, "unit": item.unit,
            "generic": item.drug.generic_name if item.drug else "",
            "available": info["available"], "mrp": str(info["mrp"] or ""),
            "schedule": item.schedule, "loose": item.sells_loose,
        })
    return JsonResponse({"results": results})


@login_required
@role_required(*PHARMACY_ROLES)
def pharmacy_bill_print(request, bill_id):
    bill = get_object_or_404(
        PharmacyBill.objects.select_related("patient", "doctor", "created_by"), id=bill_id
    )
    return render(request, "pharmacy/bill_print.html", {
        "bill": bill,
        "items": bill.items.select_related("item", "batch"),
        "gst_rows": bill.gst_summary(),
        "pharmacy": {
            "name": settings.PHARMACY_NAME,
            "drug_licence": settings.PHARMACY_DRUG_LICENCE_NO,
            "gstin": settings.PHARMACY_GSTIN,
        },
    })


@login_required
@role_required(*PHARMACY_ROLES)
@require_POST
def pharmacy_bill_cancel(request, bill_id):
    bill = get_object_or_404(PharmacyBill.objects.select_related("admission"), id=bill_id)
    reason = request.POST.get("reason", "").strip()
    if bill.is_ipd and bill.admission.status != "ADMITTED":
        messages.error(request, f"{bill.admission.ipd_no} is discharged; this issue can't be cancelled any more.")
    elif not reason:
        messages.error(request, "Enter a reason to cancel the bill.")
    else:
        cancel_bill(bill, user=request.user, reason=reason)
        messages.success(request, f"Bill {bill.bill_no} cancelled; stock returned to inventory.")
    return redirect("hms:pharmacy_bill_print", bill_id=bill.id)


@login_required
@role_required(*PHARMACY_ROLES)
def pharmacy_bills(request):
    today = timezone.localdate()
    date_from = _parse_date(request.GET.get("from"), today)
    date_to = _parse_date(request.GET.get("to"), today)
    bills = (PharmacyBill.objects.filter(created_at__date__gte=date_from, created_at__date__lte=date_to)
             .select_related("patient", "created_by", "admission"))
    q = request.GET.get("q", "").strip()
    if q:
        bills = bills.filter(Q(bill_no__icontains=q) | Q(customer_name__icontains=q) | Q(customer_mobile__icontains=q))
    total, by_mode = counter_collection(date_from, date_to)
    ipd_issued = (bills.filter(status="PAID", payment_mode="IPD")
                  .aggregate(t=Sum("net_amount"), r=Sum("returned_amount")))
    return render(request, "pharmacy/bills.html", {
        "bills": bills, "date_from": date_from, "date_to": date_to, "q": q,
        "total": total, "by_mode": by_mode,
        "ipd_issued": (ipd_issued["t"] or 0) - (ipd_issued["r"] or 0),
    })


@login_required
@role_required(*PHARMACY_ROLES)
def pharmacy_return(request, bill_id):
    """Return some (or all) units from a bill; stock goes back to the batch it came from."""
    bill = get_object_or_404(PharmacyBill.objects.select_related("patient", "admission"), id=bill_id)
    if bill.status == "CANCELLED":
        messages.error(request, "This bill is cancelled; nothing can be returned.")
        return redirect("hms:pharmacy_bill_print", bill_id=bill.id)
    if bill.is_ipd and bill.admission.status != "ADMITTED":
        messages.error(request, f"{bill.admission.ipd_no} is discharged; returns can't change its bill any more.")
        return redirect("hms:pharmacy_bill_print", bill_id=bill.id)

    if request.method == "POST":
        quantities = {k[len("ret_"):]: v for k, v in request.POST.items() if k.startswith("ret_") and v}
        try:
            ret = return_items(bill, quantities=quantities, user=request.user,
                               reason=request.POST.get("reason", "").strip())
        except ValueError as e:
            messages.error(request, str(e))
        else:
            if bill.is_ipd:
                messages.success(request, f"{ret.return_no}: ₹{ret.refund_amount} taken off {bill.admission.ipd_no}'s discharge bill")
            else:
                messages.success(request, f"{ret.return_no}: refund ₹{ret.refund_amount} ({bill.get_payment_mode_display()})")
            return redirect("hms:pharmacy_return_print", return_id=ret.id)

    return render(request, "pharmacy/return_form.html", {
        "bill": bill,
        "items": bill.items.select_related("item", "batch"),
        "factor": bill.discount_factor,
    })


@login_required
@role_required(*PHARMACY_ROLES)
def pharmacy_return_print(request, return_id):
    ret = get_object_or_404(
        PharmacyReturn.objects.select_related("bill__patient", "bill__admission", "created_by"), id=return_id
    )
    return render(request, "pharmacy/return_print.html", {
        "ret": ret,
        "items": ret.items.select_related("bill_item__item", "bill_item__batch"),
        "pharmacy": {
            "name": settings.PHARMACY_NAME,
            "drug_licence": settings.PHARMACY_DRUG_LICENCE_NO,
            "gstin": settings.PHARMACY_GSTIN,
        },
    })


@login_required
@role_required(*PHARMACY_ROLES + ("reception", "doctor", "nursing"))
def pharmacy_ipd_statement(request, admission_id):
    """Every medicine issued to an admission, with returns: backs the discharge bill's pharmacy line."""
    admission = get_object_or_404(IPDAdmission.objects.select_related("patient", "bed__ward", "doctor"), id=admission_id)
    bills = (PharmacyBill.objects.filter(admission=admission)
             .prefetch_related("items__item", "items__batch", "returns__items__bill_item__item")
             .order_by("created_at"))
    return render(request, "pharmacy/ipd_statement.html", {
        "admission": admission,
        "bills": bills,
        "total": ipd_pharmacy_total(admission),
        "pharmacy_name": settings.PHARMACY_NAME,
    })


@login_required
@role_required(*PHARMACY_ROLES)
def pharmacy_h1_register(request):
    """Schedule H1 register: every H1/X sale with patient, prescriber, drug, quantity and batch."""
    today = timezone.localdate()
    date_from = _parse_date(request.GET.get("from"), today.replace(day=1))
    date_to = _parse_date(request.GET.get("to"), today)
    entries = (
        PharmacyBillItem.objects.filter(
            item__schedule__in=("H1", "X"), bill__status="PAID",
            bill__created_at__date__gte=date_from, bill__created_at__date__lte=date_to,
        )
        .select_related("bill__patient", "bill__doctor", "item", "batch")
        .order_by("bill__created_at", "id")
    )
    return render(request, "pharmacy/h1_register.html", {
        "entries": entries, "date_from": date_from, "date_to": date_to,
        "pharmacy_name": settings.PHARMACY_NAME,
        "drug_licence": settings.PHARMACY_DRUG_LICENCE_NO,
    })


# ── Opening stock import ──────────────────────────────────────────

IMPORT_COLUMNS = [
    "Name", "Unit", "Pack Size", "Schedule (OTC/H/H1/X)", "GST %", "HSN",
    "Batch No", "Expiry (YYYY-MM-DD)", "Quantity (units)", "Purchase Price per unit",
    "MRP per unit", "Minimum Stock",
]


@login_required
@role_required(*PHARMACY_ROLES)
def pharmacy_stock_import(request):
    if request.GET.get("template"):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Opening Stock"
        ws.append(IMPORT_COLUMNS)
        ws.append(["Tab. Pantocid", "Tablet", 15, "H", 12, "3004", "PT2401", "2027-06-30", 150, 8.50, 13.20, 30])
        ws.append(["Syp. Stonil", "Bottle", 1, "OTC", 12, "3004", "ST118", "2027-01-31", 12, 95, 145, 5])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 18
        resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = 'attachment; filename="opening_stock_template.xlsx"'
        wb.save(resp)
        return resp

    results = None
    if request.method == "POST" and request.FILES.get("file"):
        results = _import_opening_stock(request.FILES["file"], request.user)
        if results["errors"]:
            messages.error(request, f"Nothing imported: fix {len(results['errors'])} row(s) and upload again.")
        else:
            messages.success(request, f"Imported {results['batches']} batch(es); {results['new_items']} new item(s) created.")
    return render(request, "pharmacy/stock_import.html", {"columns": IMPORT_COLUMNS, "results": results})


def _import_opening_stock(upload, user):
    """All-or-nothing: validate every row first, then import inside one transaction."""
    from openpyxl import load_workbook

    ws = load_workbook(upload, data_only=True).active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    drugs = {normalize_name(d.name): d for d in DrugMaster.objects.all()}
    units = {u for u, _ in InventoryItem.UNIT_CHOICES}
    schedules = {s for s, _ in InventoryItem.SCHEDULE_CHOICES}
    parsed, errors = [], []

    for n, row in enumerate(rows, start=2):
        if not row or not any(row):
            continue
        row = list(row) + [None] * (len(IMPORT_COLUMNS) - len(row))
        name, unit, pack, sched, gst, hsn, batch, expiry, qty, pp, mrp, min_stock = row[:12]
        try:
            name = str(name or "").strip()
            unit = str(unit or "Tablet").strip().title()
            sched = str(sched or "OTC").strip().upper()
            if not name:
                raise ValueError("Name is empty")
            if unit not in units:
                raise ValueError(f"Unit '{unit}' must be one of {', '.join(sorted(units))}")
            if sched not in schedules:
                raise ValueError(f"Schedule '{sched}' must be OTC, H, H1 or X")
            if hasattr(expiry, "date"):
                expiry = expiry.date()
            elif expiry:
                expiry = datetime.strptime(str(expiry).strip()[:10], "%Y-%m-%d").date()
            qty = int(qty or 0)
            mrp = Decimal(str(mrp or 0))
            if qty <= 0:
                raise ValueError("Quantity must be more than 0")
            if mrp <= 0:
                raise ValueError("MRP per unit must be more than 0")
            parsed.append(dict(
                name=name, unit=unit, pack=int(pack or 1), sched=sched,
                gst=Decimal(str(gst if gst not in (None, "") else 12)), hsn=str(hsn or "").strip(),
                batch=str(batch or "").strip(), expiry=expiry or None, qty=qty,
                pp=Decimal(str(pp or 0)), mrp=mrp, min_stock=int(min_stock or 10),
            ))
        except (ValueError, InvalidOperation, TypeError) as e:
            errors.append((n, str(e)))

    result = {"errors": errors, "batches": 0, "new_items": 0}
    if errors:
        return result
    with transaction.atomic():
        existing = {normalize_name(i.name): i for i in InventoryItem.objects.all()}
        for r in parsed:
            key = normalize_name(r["name"])
            item = existing.get(key)
            if item is None:
                item = InventoryItem.objects.create(
                    name=r["name"], category="Medicine", unit=r["unit"], pack_size=r["pack"],
                    schedule=r["sched"], gst_percent=r["gst"], hsn_code=r["hsn"],
                    minimum_stock=r["min_stock"], drug=drugs.get(key),
                )
                existing[key] = item
                result["new_items"] += 1
            receive_stock(item=item, quantity=r["qty"], batch_no=r["batch"], expiry_date=r["expiry"],
                          purchase_price=r["pp"], mrp=r["mrp"], user=user, notes="Opening stock import")
            result["batches"] += 1
    return result
